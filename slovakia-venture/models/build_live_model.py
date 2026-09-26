#!/usr/bin/env python3
"""Rebuild Addendum_A_Model.xlsx as a live, formula-driven workbook.

The source workbook (16/09/2026) contains 0 formulas in 632 cells — every
figure is a typed constant, so no assumption can be flexed. This reproduces
its structure and numbers, driven by a single Assumptions sheet.

    python3 build_live_model.py            # writes Addendum_A_Model_LIVE.xlsx
"""
from __future__ import annotations

import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = pathlib.Path(__file__).resolve().parent / "Addendum_A_Model_LIVE.xlsx"

INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # editable input
CALC_FILL = PatternFill("solid", fgColor="EAF1F8")    # calculated
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=13, color="1F3864")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EUR = '#,##0'
EUR2 = '#,##0.00'
PCT = '0.0%'
NUM2 = '0.00'


def header(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if note:
        ws["A2"] = note
        ws["A2"].font = Font(italic=True, size=9, color="595959")


def put(ws, row, col, value, *, fmt=None, fill=None, bold=False, border=True):
    c = ws.cell(row=row, column=col, value=value)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if bold:
        c.font = BOLD
    if border:
        c.border = BOX
    return c


def build() -> pathlib.Path:
    wb = Workbook()

    # ------------------------------------------------------------------ README
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 118
    lines = [
        ("Addendum A — LIVE model (rebuilt 17/09/2026)", TITLE_FONT),
        ("", None),
        ("Why this exists: the source Addendum_A_Model.xlsx (16/09/2026) contains ZERO formulas in 632 cells.", BOLD),
        ("Every figure in it is a typed constant, so no assumption can be tested. This rebuild reproduces its", None),
        ("structure and its numbers, but drives them from one Assumptions sheet.", None),
        ("", None),
        ("HOW TO USE", BOLD),
        ("Yellow cells on Assumptions are inputs. Change one and every sheet recalculates.", None),
        ("Blue cells are calculated — do not type over them.", None),
        ("", None),
        ("SHEETS", BOLD),
        ("Assumptions          every input, in one place", None),
        ("Income_Bridge_Revised  the Irish tax correction — THE key sheet. Verify with an Irish tax adviser.", None),
        ("PL_AssetLight        5-year P&L for the kitting-first business", None),
        ("Exit_Value           build-for-exit at 10% Irish CGT (Revised Entrepreneur Relief)", None),
        ("Funding_Portfolio    portfolio probability: how independent applications reach >=90%", None),
        ("Sector_Scoring_v2    the dossier's scoring sheet PLUS a Reversibility column", None),
        ("Tier_Ladder          positioning reference (descriptive, not calculated)", None),
        ("", None),
        ("WHAT CHANGED vs THE SOURCE", BOLD),
        ("1. PRSI is an input. The source hardcodes 4.2% but notes it rises to 4.35% from 10/2026.", None),
        ("2. Corporate tax applies the Slovak banded rate (10% at or below the revenue threshold, else 21%),", None),
        ("   matching Venture_Financial_Model.xlsx. The source applied a flat 21% to every year.", None),
        ("3. Exit_Value 'Years of target' is computed. The source stored 5.2 / 8.8 / 15.1 where the arithmetic", None),
        ("   gives 5.25 / 8.75 / 15.14 (one cell stored as 8.800000000000001 — hand-typed, not computed).", None),
        ("4. Sector_Scoring_v2 adds Reversibility: how much committed capital survives failure. The dossier's", None),
        ("   six criteria have no measure of capital at risk before revenue, which is why its sheet ranks the", None),
        ("   EUR600k bakery #1 while Addendum A rejects it. With Reversibility weighted, the ranking inverts.", None),
        ("5. DEPRECIATION IS CORRECTED. The source charges a flat 20,000 from 2028, i.e. 160,000/8 - but the", BOLD),
        ("   40,000 racking is not bought until 2029, so it depreciates an asset a year before purchase.", None),
        ("   Charged as incurred: 2028 = 15,000, 2029 onwards = 20,000. Effect: 2028 PBT is 19,280, not", None),
        ("   14,280. Immaterial to any gate, but a live model should be right. Set the racking capex to 0", None),
        ("   and the 2029 step disappears - which is the point of having inputs.", None),
        ("", None),
        ("Tags: (F) fact from a source · (R) reasoned estimate · (O) opinion/assumption.", None),
        ("Model numbers are (O) planning assumptions unless a source is cited.", None),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if font:
            c.font = font

    # ------------------------------------------------------------- Assumptions
    ws = wb.create_sheet("Assumptions")
    title(ws, "Assumptions — yellow cells are inputs", "Change an input and every other sheet recalculates.")
    header(ws, 4, ["Parameter", "Value", "Unit", "Source / note"], [52, 16, 12, 78])

    rows = [
        ("PERSONAL TAX (Ireland)", None, None, None),
        ("Irish income tax — higher rate", 0.40, PCT, "(F) 2026 higher rate"),
        ("USC — top band", 0.08, PCT, "(F)"),
        ("PRSI", 0.042, PCT, "(F) 4.2%; rises to 4.35% from 10/2026 — change this cell then"),
        ("Slovak dividend withholding (credited)", 0.07, PCT, "(F) treaty credit against Irish liability"),
        ("TARGET & OWNERSHIP", None, None, None),
        ("Target net income", 7500, EUR, "(F) your stated goal, per month"),
        ("Your ownership", 1.00, PCT, "(O) 100% asset-light; original plan assumed 70%"),
        ("Dividend payout ratio", 0.90, PCT, "(O) 10% retained"),
        ("COMPANY TAX (Slovakia)", None, None, None),
        ("Corporate tax — low rate", 0.10, PCT, "(R) SK CIT 10% for revenue <= threshold — confirm with accountant"),
        ("Low-rate revenue threshold", 100000, EUR, "(R)"),
        ("Corporate tax — standard rate", 0.21, PCT, "(R) 21%; 24% only above EUR5m revenue"),
        ("OPERATIONS", None, None, None),
        ("Pack price", 28, EUR2, "(R) LOAD-BEARING and UNVERIFIED. Output of the 'Cost the Pack' sprint"),
        ("Capex — container kitting line", 120000, EUR, "(F, Addendum A) 2028"),
        ("Capex — racking + WMS", 40000, EUR, "(F, Addendum A) 2029"),
        ("Depreciation period", 8, NUM2, "(O) straight line, years"),
        ("Grant released to income (per yr)", 10000, EUR, "(O) from 2028"),
        ("EXIT", None, None, None),
        ("EBITDA multiple at exit", 5, NUM2, "(O) small-company trade sale"),
        ("Irish CGT — Entrepreneur Relief rate", 0.10, PCT, "(F) Revised ER, from 01/01/2026"),
        ("Entrepreneur Relief lifetime limit", 1500000, EUR, "(F) raised from EUR1m in Budget 2026"),
        ("Irish CGT — standard rate above limit", 0.33, PCT, "(F)"),
    ]
    r = 5
    anchor = {}
    for label, val, fmt, note in rows:
        if val is None:
            c = put(ws, r, 1, label, bold=True, fill=HDR_FILL)
            c.font = Font(bold=True, color="FFFFFF")
            for col in range(2, 5):
                put(ws, r, col, None, fill=HDR_FILL)
        else:
            put(ws, r, 1, label)
            put(ws, r, 2, val, fmt=fmt, fill=INPUT_FILL, bold=True)
            put(ws, r, 3, "EUR" if fmt in (EUR, EUR2) else ("%" if fmt == PCT else ""))
            put(ws, r, 4, note)
            anchor[label] = f"Assumptions!$B${r}"
        r += 1

    A = anchor
    marg_r = r + 1
    put(ws, marg_r, 1, "Irish marginal rate (computed)", bold=True)
    put(ws, marg_r, 2, f"={A['Irish income tax — higher rate']}+{A['USC — top band']}+{A['PRSI']}",
        fmt=PCT, fill=CALC_FILL, bold=True)
    put(ws, marg_r, 4, "(F) sum of the three bands; Slovak 7% is credited, not added")
    A["MARGINAL"] = f"Assumptions!$B${marg_r}"

    # --------------------------------------------------- Income_Bridge_Revised
    ws = wb.create_sheet("Income_Bridge_Revised")
    title(ws, "Income bridge — the Irish tax correction",
          "THE key sheet of Addendum A. Verify with an Irish chartered tax adviser before acting.")
    header(ws, 4, ["Step", "Model as written (7% SK)", "Reality: Irish-resident",
                   "Irish-resident, your ownership", "Note"], [46, 22, 22, 26, 62])

    tgt = A["Target net income"]
    own = A["Your ownership"]
    pay = A["Dividend payout ratio"]
    cit = A["Corporate tax — standard rate"]
    mrg = A["MARGINAL"]
    sk = A["Slovak dividend withholding (credited)"]

    put(ws, 5, 1, "Target net income per year")
    for col, f in ((2, f"={tgt}*12"), (3, f"={tgt}*12"), (4, f"={tgt}*12")):
        put(ws, 5, col, f, fmt=EUR, fill=CALC_FILL)
    put(ws, 5, 5, "your goal x 12")

    put(ws, 6, 1, "Gross dividend needed")
    put(ws, 6, 2, f"=B5/(1-{sk})", fmt=EUR, fill=CALC_FILL)
    put(ws, 6, 3, f"=C5/(1-{mrg})", fmt=EUR, fill=CALC_FILL)
    put(ws, 6, 4, f"=D5/(1-{mrg})", fmt=EUR, fill=CALC_FILL)
    put(ws, 6, 5, "personal tax on the dividend")

    put(ws, 7, 1, "/ your ownership")
    put(ws, 7, 2, "=B6/0.7", fmt=EUR, fill=CALC_FILL)
    put(ws, 7, 3, "=C6/0.7", fmt=EUR, fill=CALC_FILL)
    put(ws, 7, 4, f"=D6/{own}", fmt=EUR, fill=CALC_FILL)
    put(ws, 7, 5, "70% in the original plan; the 4th column uses your Assumptions input")

    put(ws, 8, 1, "/ payout ratio")
    for col in (2, 3, 4):
        put(ws, 8, col, f"={get_column_letter(col)}7/{pay}", fmt=EUR, fill=CALC_FILL)
    put(ws, 8, 5, "10% retained")

    put(ws, 9, 1, "PROFIT BEFORE TAX NEEDED", bold=True)
    for col in (2, 3, 4):
        put(ws, 9, col, f"={get_column_letter(col)}8/(1-{cit})", fmt=EUR, fill=CALC_FILL, bold=True)
    put(ws, 9, 5, "at the Slovak standard corporate rate")

    for i, m in enumerate((0.08, 0.12, 0.20)):
        rr = 10 + i
        put(ws, rr, 1, f"Revenue needed at {m:.0%} pre-tax margin")
        for col in (2, 3, 4):
            put(ws, rr, col, f"={get_column_letter(col)}9/{m}", fmt=EUR, fill=CALC_FILL)
        put(ws, rr, 5, "(R)")

    put(ws, 14, 1, "Irish marginal rate build-up", bold=True)
    for i, (lbl, ref) in enumerate((
        ("Income tax", A["Irish income tax — higher rate"]),
        ("USC", A["USC — top band"]),
        ("PRSI", A["PRSI"]),
    )):
        put(ws, 15 + i, 1, lbl)
        put(ws, 15 + i, 2, f"={ref}", fmt=PCT, fill=CALC_FILL)
    put(ws, 18, 1, "Total marginal rate", bold=True)
    put(ws, 18, 2, f"={mrg}", fmt=PCT, fill=CALC_FILL, bold=True)
    put(ws, 18, 5, "credit given for the 7% Slovak withholding (F, treaty)")

    # ----------------------------------------------------------- PL_AssetLight
    ws = wb.create_sheet("PL_AssetLight")
    title(ws, "P&L — asset-light (kitting-first) business",
          "Revenue lines are inputs (yellow). Everything below Gross profit is calculated.")
    years = [2027, 2028, 2029, 2030, 2031]
    header(ws, 4, ["Line"] + [str(y) for y in years], [46, 15, 15, 15, 15, 15])

    rev = {
        "Revenue — brokerage (Tier 0)": [150000, 200000, 150000, 120000, 100000],
        "Revenue — kitting (Tier 1)": [0, 380000, 950000, 1600000, 2100000],
        "Revenue — reserve service (Tier 2)": [0, 40000, 120000, 260000, 420000],
        "Revenue — software (LEANTA)": [12000, 48000, 96000, 144000, 180000],
    }
    r = 5
    for label, vals in rev.items():
        put(ws, r, 1, label)
        for i, v in enumerate(vals):
            put(ws, r, 2 + i, v, fmt=EUR, fill=INPUT_FILL)
        r += 1
    rev_first, rev_last = 5, r - 1

    put(ws, r, 1, "TOTAL REVENUE", bold=True)
    for i in range(5):
        col = get_column_letter(2 + i)
        put(ws, r, 2 + i, f"=SUM({col}{rev_first}:{col}{rev_last})", fmt=EUR, fill=CALC_FILL, bold=True)
    total_r = r; r += 1

    put(ws, r, 1, "Gross margin %")
    for i, v in enumerate([0.13, 0.21, 0.25, 0.27, 0.28]):
        put(ws, r, 2 + i, v, fmt=PCT, fill=INPUT_FILL)
    gm_r = r; r += 1

    put(ws, r, 1, "Gross profit")
    for i in range(5):
        col = get_column_letter(2 + i)
        put(ws, r, 2 + i, f"={col}{total_r}*{col}{gm_r}", fmt=EUR, fill=CALC_FILL)
    gp_r = r; r += 1

    put(ws, r, 1, "Fixed overheads")
    for i, v in enumerate([28000, 110000, 190000, 280000, 350000]):
        put(ws, r, 2 + i, v, fmt=EUR, fill=INPUT_FILL)
    fo_r = r; r += 1

    put(ws, r, 1, "EBITDA", bold=True)
    for i in range(5):
        col = get_column_letter(2 + i)
        put(ws, r, 2 + i, f"={col}{gp_r}-{col}{fo_r}", fmt=EUR, fill=CALC_FILL, bold=True)
    ebitda_r = r; r += 1

    # Depreciation charged on capex ACTUALLY INCURRED to date: the kitting line
    # from 2028, the racking/WMS only from 2029. The source workbook charges a
    # flat 160,000/8 from 2028, i.e. it depreciates the racking a year before
    # it is bought. See README item 5.
    put(ws, r, 1, "Depreciation")
    line, rack, per = (A["Capex — container kitting line"],
                       A["Capex — racking + WMS"], A["Depreciation period"])
    for i, y in enumerate(years):
        if y < 2028:
            f = "0"
        elif y < 2029:
            f = f"={line}/{per}"
        else:
            f = f"=({line}+{rack})/{per}"
        put(ws, r, 2 + i, f, fmt=EUR, fill=CALC_FILL)
    dep_r = r; r += 1

    put(ws, r, 1, "Grant released to income")
    for i, y in enumerate(years):
        f = "0" if y < 2028 else f"={A['Grant released to income (per yr)']}"
        put(ws, r, 2 + i, f, fmt=EUR, fill=CALC_FILL)
    gr_r = r; r += 1

    put(ws, r, 1, "Interest")
    for i, v in enumerate([0, 6000, 6000, 5000, 4000]):
        put(ws, r, 2 + i, v, fmt=EUR, fill=INPUT_FILL)
    int_r = r; r += 1

    put(ws, r, 1, "PROFIT BEFORE TAX", bold=True)
    for i in range(5):
        col = get_column_letter(2 + i)
        put(ws, r, 2 + i,
            f"={col}{ebitda_r}-{col}{dep_r}+{col}{gr_r}-{col}{int_r}",
            fmt=EUR, fill=CALC_FILL, bold=True)
    pbt_r = r; r += 1

    put(ws, r, 1, "Corporate tax")
    for i in range(5):
        col = get_column_letter(2 + i)
        put(ws, r, 2 + i,
            f"=IF({col}{pbt_r}<=0,0,{col}{pbt_r}*IF({col}{total_r}<="
            f"{A['Low-rate revenue threshold']},{A['Corporate tax — low rate']},"
            f"{A['Corporate tax — standard rate']}))",
            fmt=EUR, fill=CALC_FILL)
    tax_r = r; r += 1

    put(ws, r, 1, "Net profit", bold=True)
    for i in range(5):
        col = get_column_letter(2 + i)
        put(ws, r, 2 + i, f"={col}{pbt_r}-{col}{tax_r}", fmt=EUR, fill=CALC_FILL, bold=True)
    np_r = r; r += 1

    put(ws, r, 1, "Your net income / month (Irish-resident)", bold=True)
    for i in range(5):
        col = get_column_letter(2 + i)
        put(ws, r, 2 + i,
            f"=MAX(0,{col}{np_r})*{pay}*{own}*(1-{mrg})/12",
            fmt=EUR, fill=CALC_FILL, bold=True)
    inc_r = r; r += 1

    put(ws, r, 1, "Gap to target")
    for i in range(5):
        col = get_column_letter(2 + i)
        put(ws, r, 2 + i, f"={col}{inc_r}-{tgt}", fmt=EUR, fill=CALC_FILL)
    r += 1
    put(ws, r, 1, "Target met?")
    for i in range(5):
        col = get_column_letter(2 + i)
        put(ws, r, 2 + i, f'=IF({col}{inc_r}>={tgt},"Yes","No")', fill=CALC_FILL)
    r += 2

    for label, text in (
        ("GATE 0 — 30/06/2027", ">=EUR50k Tier 0 revenue + 1 state reference; else STOP (<EUR8k spent)"),
        ("GATE 1 — 31/12/2027", ">=2 co-packing agreements + >=EUR300k contracted demand + grant decision; else stay a broker"),
        ("GATE 2 — 31/12/2029", "EBITDA >= EUR100k + reserve contract signed; else sell the book, keep the software"),
    ):
        put(ws, r, 1, label, bold=True)
        put(ws, r, 2, text)
        r += 1

    # ------------------------------------------------------------- Exit_Value
    ws = wb.create_sheet("Exit_Value")
    title(ws, "Exit value at 10% Irish CGT (Revised Entrepreneur Relief)",
          "(O) Compare: the same money taken as dividends while Irish-resident is taxed at the marginal rate.")
    header(ws, 4, ["EBITDA at exit", "Multiple", "Enterprise value", "Ownership",
                   "Your share", "CGT", "Net to you", "Years of target"],
           [16, 10, 18, 12, 16, 14, 16, 14])
    r = 5
    for ebitda in (150000, 250000, 400000, "=PL_AssetLight!F" + str(ebitda_r)):
        for ownership in (0.70, 1.00):
            put(ws, r, 1, ebitda, fmt=EUR, fill=INPUT_FILL if isinstance(ebitda, int) else CALC_FILL)
            put(ws, r, 2, f"={A['EBITDA multiple at exit']}", fmt=NUM2, fill=CALC_FILL)
            put(ws, r, 3, f"=A{r}*B{r}", fmt=EUR, fill=CALC_FILL)
            put(ws, r, 4, ownership, fmt=PCT, fill=INPUT_FILL)
            put(ws, r, 5, f"=C{r}*D{r}", fmt=EUR, fill=CALC_FILL)
            lim = A["Entrepreneur Relief lifetime limit"]
            put(ws, r, 6,
                f"=MIN(E{r},{lim})*{A['Irish CGT — Entrepreneur Relief rate']}"
                f"+MAX(0,E{r}-{lim})*{A['Irish CGT — standard rate above limit']}",
                fmt=EUR, fill=CALC_FILL)
            put(ws, r, 7, f"=E{r}-F{r}", fmt=EUR, fill=CALC_FILL, bold=True)
            put(ws, r, 8, f"=G{r}/({tgt}*12)", fmt=NUM2, fill=CALC_FILL)
            r += 1
    put(ws, r + 1, 1,
        "(F) Revised Entrepreneur Relief: 10% CGT, lifetime limit EUR1.5m from 01/01/2026; >=5% ordinary "
        "shares; director/employee >=50% of working time for 3 of the previous 5 years.")

    # -------------------------------------------------------- Funding_Portfolio
    ws = wb.create_sheet("Funding_Portfolio")
    title(ws, "Funding portfolio — probability that at least one application succeeds",
          "(R) Probabilities are judgement, not published statistics. Edit them; the maths follows.")
    header(ws, 4, ["#", "Instrument", "Type / size", "Use", "Success (R)",
                   "Include?", "(1-p) if included"],
           [5, 42, 34, 30, 12, 10, 16])
    instruments = [
        ("LEO Feasibility grant (IE)", "grant, ~EUR15k, 50%", "Market study", 0.60, 1),
        ("LEO Priming grant (IE)", "50% or EUR150,000 max (F)", "Irish entity first-year costs", 0.45, 1),
        ("Enterprise Ireland / IDA defence", "grants/equity/loans (F, 2026 Bill)", "Export & capability", 0.30, 1),
        ("Regionalny prispevok (MIRRI, RS)", "share of EUR13.9m; deadline 05/10/2026 (F)", "Kitting line, jobs", 0.45, 1),
        ("Just Transition Fund (BB region)", "MSME grants to EUR800,000 (F)", "Productive investment", 0.40, 1),
        ("PPA intervention 73.7", "EUR35m call (F)", "Tier 3 line", 0.35, 1),
        ("LEADER / CLLD via MAS", "EUR10-100k (R)", "Small equipment", 0.50, 1),
        ("Interreg SK-HU small project fund", "small cross-border (R)", "Preparedness pilot", 0.40, 1),
        ("Visegrad Fund", "EUR25-35k; small grants <=EUR6k (F)", "Education project only", 0.40, 1),
        ("SBA / NHF microloan", "EUR2,500-50,000; 1.19-9.03% (F)", "Equipment / working capital", 0.68, 1),
        ("SZRB INVESTaktiv (EIF guarantee)", "investment to 20 yrs (F)", "Tier 3 plant", 0.60, 1),
        ("SIH / MH SR NDF III", "to EUR2.8m, 80% guarantee (F)", "Tier 3 plant", 0.60, 1),
        ("Robotisation vouchers 2026", "EUR3-30k at 85% (F)", "Automation", 0.55, 1),
        ("EDIP", "EUR1.5bn WP 2026-27 (F)", "Later, consortium only", 0.10, 1),
    ]
    r = 5
    for i, (name, typ, use, p, inc) in enumerate(instruments, start=1):
        put(ws, r, 1, i)
        put(ws, r, 2, name)
        put(ws, r, 3, typ)
        put(ws, r, 4, use)
        put(ws, r, 5, p, fmt=PCT, fill=INPUT_FILL)
        put(ws, r, 6, inc, fill=INPUT_FILL)
        put(ws, r, 7, f"=IF(F{r}=1,1-E{r},1)", fmt=NUM2, fill=CALC_FILL)
        r += 1
    first, last = 5, r - 1
    r += 1
    put(ws, r, 2, "P(at least one succeeds), included rows", bold=True)
    # Plain PRODUCT over a helper column, not an array formula: PRODUCT(IF(...))
    # needs Ctrl+Shift+Enter in pre-dynamic-array Excel and silently returns the
    # wrong number without it.
    put(ws, r, 5, f"=1-PRODUCT($G${first}:$G${last})", fmt=PCT, fill=CALC_FILL, bold=True)
    r += 1
    put(ws, r, 2, "Set Include? to 0/1 to test a smaller portfolio", border=False)
    r += 1
    put(ws, r, 2, "Rule (O): submit >=3 INDEPENDENT applications, different systems, different calendars.")
    r += 1
    put(ws, r, 2, "Three applications to the same ministry are one application wearing three hats.")

    # ------------------------------------------------------- Sector_Scoring_v2
    ws = wb.create_sheet("Sector_Scoring_v2")
    title(ws, "Sector scoring v2 — with a Reversibility criterion",
          "The dossier's sheet had no measure of capital at risk before revenue. That omission is why it "
          "ranks the EUR600k bakery #1 while Addendum A rejects it.")
    crit = ["Fit with skills", "Demand security", "Low capital need",
            "Grant eligibility", "Margin", "Speed to cash", "Reversibility"]
    header(ws, 4, ["Option"] + crit + ["Weighted score", "Rank"],
           [36] + [13] * 7 + [15, 8])

    put(ws, 5, 1, "Weight (edit)", bold=True)
    for i, w in enumerate([0.20, 0.16, 0.12, 0.12, 0.12, 0.08, 0.20]):
        put(ws, 5, 2 + i, w, fmt=PCT, fill=INPUT_FILL, bold=True)
    put(ws, 5, 9, "=SUM(B5:H5)", fmt=PCT, fill=CALC_FILL, bold=True)
    put(ws, 5, 10, '=IF(ABS(I5-1)<0.0001,"OK","WEIGHTS != 100%")')

    options = [
        ("Tier 0 — broker / agent", [4, 4, 5, 2, 2, 5, 5]),
        ("Tier 1 — kitter / assembler", [4, 4, 4, 4, 3, 3, 4]),
        ("Tier 2 — reserve service", [4, 5, 4, 3, 4, 2, 3]),
        ("Shelf-stable bakery (dossier rec.)", [5, 4, 2, 5, 3, 3, 1]),
        ("Lubricants: distribution -> blending", [2, 4, 4, 3, 4, 4, 3]),
        ("Wooden packaging / pallets / crates", [3, 4, 3, 4, 2, 3, 2]),
    ]
    r = 6
    for name, scores in options:
        put(ws, r, 1, name)
        for i, s in enumerate(scores):
            put(ws, r, 2 + i, s, fill=INPUT_FILL)
        put(ws, r, 9, f"=SUMPRODUCT($B$5:$H$5,B{r}:H{r})", fmt=NUM2, fill=CALC_FILL, bold=True)
        put(ws, r, 10, f"=RANK(I{r},$I$6:$I${5 + len(options)})", fill=CALC_FILL)
        r += 1
    r += 1
    for note in (
        "(F) Original dossier weights: Fit 25% · Demand 20% · Low capital 15% · Grant 15% · Margin 15% · Speed 10%.",
        "(O) Reversibility added at 20%, the other six scaled by 0.8 so the weights still sum to 100%.",
        "(O) Reversibility = how much committed capital survives if the venture fails. A container line resells;",
        "    a hall fit-out and a certification do not.",
        "(R) Scores are judgement. Replace them with yours — that is the point of this sheet.",
    ):
        put(ws, r, 1, note, border=False)
        r += 1

    # -------------------------------------------------------------- Tier_Ladder
    ws = wb.create_sheet("Tier_Ladder")
    title(ws, "Tier ladder — positioning from cheapest to most expensive", "Descriptive reference (F, Addendum A).")
    header(ws, 4, ["Tier", "What you sell", "Capex (EUR)", "Gross margin",
                   "Time to first cash", "Win probability (R)", "Verdict"],
           [26, 46, 20, 16, 20, 20, 30])
    ladder = [
        ("0 — Broker/agent", "Resell certified rations & kits into SK tenders, back-to-back",
         "2,000-5,000", "6-12%", "60-90 days", "High", "Start here Q4 2026"),
        ("1 — Kitter/assembler", "Assemble ration packs & 72h civil kits under own HACCP",
         "120,000-180,000", "18-26%", "9-14 months", "Medium-high", "Core business 2027-28"),
        ("2 — Reserve service", "Hold & rotate state/municipal stock for an annual fee",
         "40,000-90,000", "25-40% on fee", "12-24 months", "Medium", "The annuity — highest value"),
        ("3 — Own production", "Shelf-stable bakery / retort / lyo line (original plan)",
         "450,000-700,000", "30-38%", "24-36 months", "Medium", "Only after Tier 1 contracts"),
        ("4 — Deployable modules", "Containerised field bakery/kitchen/packing, sold or leased",
         "150,000-400,000/unit", "25-35%", "18-30 months", "Low-medium", "Export & NSPA upside"),
    ]
    for i, row in enumerate(ladder):
        for j, v in enumerate(row):
            put(ws, 5 + i, 1 + j, v)

    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    p = build()
    print(f"written: {p}")
